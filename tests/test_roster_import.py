from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

import scripts.import_monitor_users as importer
from scripts.import_monitor_users import load_roster_plan


@pytest.fixture
def roster_workbook(tmp_path: Path) -> Path:
    path = tmp_path / "OurA-Navi_userlist.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "OurA-Naviユーザー管理"
    sheet.append([
        "エリア",
        "勤務地",
        "社員名",
        "社員メールアドレス",
        "役割",
        "テルモMR経歴",
        "備考",
    ])
    sheet.append(["北海道東北", "札幌", "MR利用者", "mr@example.com", "本社MR", "10年", "MR"])
    sheet.append(["本社", "虎ノ門", "HC利用者", "hc@example.com", "本部メンバー", "", "本社（ヘルスケア）"])
    sheet.append(["首都圏A", "東京", "DM本社利用者", "dm@example.com", "本部メンバー", "", "本社（DM）"])
    sheet.append(["本社", "虎ノ門", "管理者", "admin@example.com", "本部メンバー", "", "システム管理者"])
    workbook.save(path)
    return path


def test_roster_import_plan_derives_summary_from_exact_roles_and_user_map_from_departments(roster_workbook: Path) -> None:
    plan = load_roster_plan(roster_workbook)
    assert len(plan.users) == 4
    assert plan.scope_counts == {"global": 1, "user_map": 3, "management": 4}
    assert plan.department_counts == {
        "MR(DM)": 1,
        "ヘルスケア本社": 1,
        "DM本社": 1,
        "管理者": 1,
    }
    assert all(item["user_id"] == "" for item in plan.users)
    assert all("user_key" not in item for item in plan.users)


def test_toranomon_and_tokyo_business_area_remain_distinct(roster_workbook: Path) -> None:
    plan = load_roster_plan(roster_workbook)
    headquarters = [item for item in plan.users if item["area_key"] == "本社・虎ノ門"]
    tokyo_business = [item for item in plan.users if item["area_key"] == "首都圏A"]
    assert len(headquarters) == 2
    assert len(tokyo_business) == 1


def test_bootstrap_apply_resumes_only_an_exact_partial_plan(
    roster_workbook: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    plan = load_roster_plan(roster_workbook)
    stored = [dict(plan.users[0])]
    planned_by_email = {item["email"]: item for item in plan.users}

    class Directory:
        @staticmethod
        def list_users(*, include_inactive: bool = True):
            assert include_inactive is True
            return [dict(item) for item in stored]

    class Manager:
        def __init__(self, *, directory):
            assert directory is not None

        @staticmethod
        def create_user(payload, *, actor: str):
            assert actor == "admin@example.com"
            stored.append(dict(planned_by_email[payload.email]))

    credential = roster_workbook.parent / "approved-credential.json"
    credential.write_text("{}", encoding="utf-8")
    credential.chmod(0o600)
    monkeypatch.setattr(
        importer,
        "get_settings",
        lambda: SimpleNamespace(
            monitor_firestore_database="lcs-user-data",
            monitor_project_id="test-project",
        ),
    )
    monkeypatch.setattr(
        importer.service_account.Credentials,
        "from_service_account_file",
        lambda _path: object(),
    )
    monkeypatch.setattr(importer.firestore, "Client", lambda **_kwargs: object())
    monkeypatch.setattr(
        importer,
        "UserDirectoryRepository",
        lambda _settings, *, client: Directory(),
    )
    monkeypatch.setattr(importer, "UserManagementService", Manager)

    receipt = importer._apply_plan(
        plan,
        actor="admin@example.com",
        credential_file=str(credential),
    )

    assert receipt == {
        "appliedUsers": 4,
        "createdUsers": 3,
        "resumedFromUsers": 1,
        "readbackMatched": True,
    }


def test_bootstrap_apply_refuses_to_overwrite_a_changed_existing_row(
    roster_workbook: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    plan = load_roster_plan(roster_workbook)
    changed = {**plan.users[0], "role": "別の役割"}

    class Directory:
        @staticmethod
        def list_users(*, include_inactive: bool = True):
            return [dict(changed)]

    credential = roster_workbook.parent / "approved-credential.json"
    credential.write_text("{}", encoding="utf-8")
    credential.chmod(0o600)
    monkeypatch.setattr(
        importer,
        "get_settings",
        lambda: SimpleNamespace(
            monitor_firestore_database="lcs-user-data",
            monitor_project_id="test-project",
        ),
    )
    monkeypatch.setattr(
        importer.service_account.Credentials,
        "from_service_account_file",
        lambda _path: object(),
    )
    monkeypatch.setattr(importer.firestore, "Client", lambda **_kwargs: object())
    monkeypatch.setattr(
        importer,
        "UserDirectoryRepository",
        lambda _settings, *, client: Directory(),
    )
    monkeypatch.setattr(
        importer,
        "UserManagementService",
        lambda **_kwargs: pytest.fail("changed bootstrap rows must be rejected before writes"),
    )

    with pytest.raises(RuntimeError, match="existing rows differ"):
        importer._apply_plan(
            plan,
            actor="admin@example.com",
            credential_file=str(credential),
        )


@pytest.fixture
def hcs_workbook(tmp_path: Path) -> Path:
    path = tmp_path / "userlist.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append([*importer.EXPECTED_HEADERS, "役割", "HCSチーム"])
    sheet.append(["東海北陸", "静岡", "移籍 利用者", "transfer@example.com", "社員MR", "20年以上", "MR（DM）", None, "ー"])
    sheet.append(["東海北陸", "静岡", "移籍 利用者", "transfer@example.com", "社員MR", "20年以上", "MR（HCS）", None, "HCS静岡ﾏﾙﾁ"])
    sheet.append(["北海道東北", "札幌", "新規 HCS", "hcs@example.com", "社員MR", "5年未満", "MR（HCS）", None, ""])
    sheet.append(["関西", "大阪", "DM 利用者", "dm@example.com", "コントラクトMR", "10年", "MR（DM）", None, "ー"])
    sheet.append(["東京", "虎ノ門", "本社 利用者", "hq@example.com", "本部メンバー", "", "本社（DM）", None, "ー"])
    workbook.save(path)
    return path


def test_new_workbook_requires_explicit_duplicate_choice_and_preserves_hcs_fields(hcs_workbook: Path) -> None:
    with pytest.raises(ValueError, match="duplicate email.*rows 2, 3"):
        load_roster_plan(hcs_workbook)
    plan = load_roster_plan(hcs_workbook, duplicate_departments={"transfer@example.com": "MR(HCS)"})
    assert plan.source_rows == 5
    assert plan.resolved_duplicate_rows == [2]
    assert plan.normalized_headquarters_rows == [6]
    assert plan.scope_counts == {"global": 3, "user_map": 4, "management": 4}
    assert plan.department_counts == {"MR(HCS)": 2, "MR(DM)": 1, "DM本社": 1}
    transferred = next(item for item in plan.users if item["email"] == "transfer@example.com")
    assert transferred["department"] == "MR(HCS)"
    assert transferred["team"] == "HCS静岡マルチ"
    assert transferred["mr_experience"] == "20年以上"
    assert next(item for item in plan.users if item["email"] == "hcs@example.com")["team"] == ""
    assert next(item for item in plan.users if item["email"] == "hq@example.com")["area_key"] == "本社・虎ノ門"


def test_duplicate_resolution_must_select_exactly_one_source_record(hcs_workbook: Path) -> None:
    with pytest.raises(ValueError, match="select one exact department"):
        load_roster_plan(hcs_workbook, duplicate_departments={"transfer@example.com": "管理者"})
    with pytest.raises(ValueError, match="did not match"):
        load_roster_plan(hcs_workbook, duplicate_departments={"transfer@example.com": "MR(HCS)", "hcs@example.com": "MR(HCS)"})


@pytest.fixture
def current_multisheet_workbook(tmp_path: Path) -> Path:
    """Synthetic October roster shape; never embed production identities in tests."""
    path = tmp_path / "user_list.xlsx"
    workbook = Workbook()
    workbook.active.title = "開発"
    workbook.active.append(["開発用の別表", "このシートを名簿として読まない"])
    workbook.create_sheet("選択肢マスタ").append(["役割", "部門"])
    sheet = workbook.create_sheet("OurA-Naviユーザー管理")
    sheet.append([*importer.EXPECTED_HEADERS, "役割", "HCSチーム"])
    for department, count in (("DM", 64), ("HCS", 74)):
        for index in range(count):
            area, workplace = "関西", "大　阪"
            if index == 0:
                area, workplace = "九州", "沖　縄"
            elif department == "HCS" and index <= 4:
                area, workplace = "北海道東北", "新　潟"
            sheet.append([
                area, workplace, f"合成{department}{index}",
                f"{department.lower()}{index}@example.com", "社員MR", "5年未満",
                f"MR（{department}）", "アンバサダー",
                "HCS新潟医療器･医薬品" if department == "HCS" else "ー",
            ])
    for group, count, remark in (
        ("dm-hq", 13, "本社（DM）"),
        ("healthcare-hq", 8, "本社（ヘルスケア）"),
        ("admin", 3, "システム管理者"),
    ):
        for index in range(count):
            sheet.append([
                "東京", "虎ノ門", f"合成本部{group}{index}", f"{group}{index}@example.com",
                "本部メンバー", "ー", remark, "DM・HRJ戦略ﾘｰﾀﾞｰ", "ー",
            ])
    geography = workbook.create_sheet("DMエリア追加 支店別対象管轄エリア ")
    geography.append([None, "■支店別対象管轄エリア"])
    geography.append([None, "支店名", "県名", "DMエリア区分"])
    geography.append([None, "新潟", "新潟県", "関東A"])
    workbook.save(path)
    return path


def test_current_multisheet_roster_uses_existing_fields_and_ignores_extra_role(
    current_multisheet_workbook: Path,
) -> None:
    plan = load_roster_plan(current_multisheet_workbook)
    assert plan.source_rows == 162
    assert plan.resolved_duplicate_rows == []
    assert plan.scope_counts == {"global": 138, "user_map": 159, "management": 162}
    assert plan.department_counts == {
        "MR(DM)": 64, "MR(HCS)": 74, "DM本社": 13, "ヘルスケア本社": 8, "管理者": 3,
    }
    assert len(plan.normalized_headquarters_rows) == 24
    assert all(
        item["area_key"] == "本社・虎ノ門" and item["mr_experience"] == "-"
        for item in plan.users if item["role"] == "本部メンバー"
    )
    assert all(item["team"] == "" for item in plan.users if item["department"] != "MR(HCS)")
    assert all(
        item["team"] == "HCS新潟医療器・医薬品" and item["mr_experience"] == "5年未満"
        for item in plan.users if item["department"] == "MR(HCS)"
    )

    # Changing every H value must not change any imported field or scope.
    workbook = load_workbook(current_multisheet_workbook)
    for row in workbook["OurA-Naviユーザー管理"].iter_rows(min_row=2, min_col=8, max_col=8):
        row[0].value = "名簿に保存しない追加職務"
    workbook.save(current_multisheet_workbook)
    workbook.close()
    assert load_roster_plan(current_multisheet_workbook) == plan


def test_roster_area_is_not_inferred_from_workplace_or_map_coverage(
    current_multisheet_workbook: Path,
) -> None:
    plan = load_roster_plan(current_multisheet_workbook)
    okinawa = [item for item in plan.users if item["workplace"] == "沖 縄"]
    assert len(okinawa) == 2
    assert {item["department"] for item in okinawa} == {"MR(DM)", "MR(HCS)"}
    assert all(item["area"] == item["area_key"] == "九州" for item in okinawa)
    niigata_hcs = [item for item in plan.users if item["workplace"] == "新 潟"]
    assert len(niigata_hcs) == 4
    assert all(item["area"] == item["area_key"] == "北海道東北" for item in niigata_hcs)


class SyncDirectory:
    def __init__(self, users):
        self.users = {item["roster_id"]: dict(item) for item in users}
        self.audit = []

    def list_users(self, *, include_inactive=True):
        return [dict(item) for item in self.users.values()]

    def get_user(self, roster_id):
        return dict(self.users[roster_id]) if roster_id in self.users else None

    def find_user_by_email(self, email):
        return next((dict(item) for item in self.users.values() if item["email"] == email), None)

    def list_labels(self, *, include_inactive=True):
        return [{"label_id": "existing-label", "is_active": False}]

    def put_user_and_change(self, user, change, *, expected_updated_at=""):
        current = self.users.get(user["roster_id"], {})
        assert importer.revision_text(current.get("updated_at")) == expected_updated_at
        self.users[user["roster_id"]] = dict(user)
        self.audit.append(change)
        return dict(user)


def test_sync_keeps_complete_roster_without_login_or_usage_records(
    current_multisheet_workbook: Path,
) -> None:
    plan = load_roster_plan(current_multisheet_workbook)
    directory = SyncDirectory([])
    diff = importer._sync_diff(plan, directory.list_users())
    receipt = importer._apply_sync(
        plan, directory=directory, actor="admin@example.com",
        expected_sync_digest=diff["summary"]["syncDigest"],
    )
    assert receipt["readbackMatched"] is True
    assert len(directory.users) == 162
    assert all(item["user_id"] == item["chat_user_id"] == "" for item in directory.users.values())
    assert all(item["is_active"] is True for item in directory.users.values())
    repeated = importer._sync_diff(plan, directory.list_users())["summary"]
    assert repeated["createUsers"] == repeated["updateUsers"] == 0


def test_sync_updates_stable_identity_preserves_history_and_is_resumable(hcs_workbook: Path) -> None:
    plan = load_roster_plan(hcs_workbook, duplicate_departments={"transfer@example.com": "MR(HCS)"})
    source = next(item for item in plan.users if item["email"] == "transfer@example.com")
    current = {
        **source, "roster_id": "existing-stable-id", "department": "DM専任", "role": "本社MR", "team": "",
        "user_id": "bound-subject", "chat_user_id": "bound-chat", "identity_bound_at": "binding-time",
        "label_ids": ["existing-label"], "is_active": False,
        "created_at": "creation-time", "updated_at": "2026-09-01T00:00:00+00:00",
    }
    outside = {**current, "roster_id": "outside-id", "email": "outside@example.com", "user_id": "outside-subject", "chat_user_id": "outside-chat"}
    directory = SyncDirectory([current, outside])
    diff = importer._sync_diff(plan, directory.list_users())
    assert diff["summary"]["createUsers"] == 3
    assert diff["summary"]["updateUsers"] == 1
    assert diff["summary"]["outsideWorkbookUsers"] == 1
    receipt = importer._apply_sync(plan, directory=directory, actor="admin@example.com", expected_sync_digest=diff["summary"]["syncDigest"])
    assert receipt["readbackMatched"] is True
    actual = directory.users["existing-stable-id"]
    assert actual["department"] == "MR(HCS)"
    assert actual["role"] == "社員MR"
    assert actual["team"] == "HCS静岡マルチ"
    for key in (*importer._PRESERVED_FIELDS, *importer._IDENTITY_FIELDS):
        assert actual[key] == current[key]
    assert directory.users["outside-id"] == outside
    assert source["roster_id"] not in directory.users
    assert len(directory.audit) == 4
    repeated = importer._sync_diff(plan, directory.list_users())
    assert repeated["summary"]["createUsers"] == 0
    assert repeated["summary"]["updateUsers"] == 0
    importer._apply_sync(plan, directory=directory, actor="admin@example.com", expected_sync_digest=repeated["summary"]["syncDigest"])
    assert len(directory.audit) == 4


def test_sync_refuses_changed_plan_before_any_writes(hcs_workbook: Path) -> None:
    plan = load_roster_plan(hcs_workbook, duplicate_departments={"transfer@example.com": "MR(HCS)"})
    directory = SyncDirectory([])
    digest = importer._sync_diff(plan, directory.list_users())["summary"]["syncDigest"]
    directory.users["concurrent"] = {**plan.users[0], "roster_id": "concurrent"}
    with pytest.raises(RuntimeError, match="sync plan changed"):
        importer._apply_sync(plan, directory=directory, actor="admin@example.com", expected_sync_digest=digest)
    assert directory.audit == []


def test_sync_refuses_ambiguous_current_email_before_any_writes(hcs_workbook: Path) -> None:
    plan = load_roster_plan(hcs_workbook, duplicate_departments={"transfer@example.com": "MR(HCS)"})
    current = plan.users[0]
    with pytest.raises(RuntimeError, match="ambiguous identity"):
        importer._sync_diff(plan, [current, {**current, "roster_id": "another-id"}])


def test_sync_refuses_duplicate_bound_identity_before_any_writes(hcs_workbook: Path) -> None:
    plan = load_roster_plan(hcs_workbook, duplicate_departments={"transfer@example.com": "MR(HCS)"})
    with pytest.raises(RuntimeError, match="duplicate_identity"):
        importer._sync_diff(plan, [{**item, "user_id": "same-subject"} for item in plan.users[:2]])


def test_sync_repairs_derived_area_key(hcs_workbook: Path) -> None:
    plan = load_roster_plan(hcs_workbook, duplicate_departments={"transfer@example.com": "MR(HCS)"})
    directory = SyncDirectory([{**item, "area_key": "wrong-region"} for item in plan.users])
    diff = importer._sync_diff(plan, directory.list_users())
    assert diff["summary"]["updateUsers"] == 4
    importer._apply_sync(plan, directory=directory, actor="admin@example.com", expected_sync_digest=diff["summary"]["syncDigest"])
    assert all(directory.users[item["roster_id"]]["area_key"] == item["area_key"] for item in plan.users)


def test_sync_refuses_missing_label_before_any_writes(hcs_workbook: Path) -> None:
    plan = load_roster_plan(hcs_workbook, duplicate_departments={"transfer@example.com": "MR(HCS)"})
    directory = SyncDirectory([{**plan.users[0], "team": "old", "label_ids": ["missing-label"]}])
    digest = importer._sync_diff(plan, directory.list_users())["summary"]["syncDigest"]
    with pytest.raises(RuntimeError, match="missing labels"):
        importer._apply_sync(plan, directory=directory, actor="admin@example.com", expected_sync_digest=digest)
    assert directory.audit == []
